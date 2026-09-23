import io

import openpyxl
import pytest
from httpx import AsyncClient

from tests.conftest import REPO_DIR, bearer, login

pytestmark = pytest.mark.integration

PROJECTS = "/api/v1/projects"
DATASET = REPO_DIR / "case" / "dataset" / "Датасеты_хакатон.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _blank(client: AsyncClient) -> tuple[str, dict[str, str]]:
    headers = bearer((await login(client, "user@robomera.demo"))["access"])
    created = await client.post(
        PROJECTS, json={"name": "Импорт", "object_type": "warehouse"}, headers=headers
    )
    return created.json()["id"], headers


async def test_template_download_is_xlsx(client: AsyncClient) -> None:
    response = await client.get("/api/v1/object-types/warehouse/template.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Параметры" in workbook.sheetnames


async def test_filled_template_roundtrip(client: AsyncClient) -> None:
    project_id, headers = await _blank(client)
    template = (await client.get("/api/v1/object-types/warehouse/template.xlsx")).content
    workbook = openpyxl.load_workbook(io.BytesIO(template))
    sheet = workbook["Параметры"]
    header = [cell.value for cell in sheet[1]]
    key_col, value_col = header.index("Ключ") + 1, header.index("Значение") + 1
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, key_col).value == "area_m2":
            sheet.cell(row, value_col).value = 7000
    buffer = io.BytesIO()
    workbook.save(buffer)

    parsed = await client.post(
        f"{PROJECTS}/{project_id}/import",
        files={"file": ("my.xlsx", buffer.getvalue(), XLSX)},
        headers=headers,
    )
    assert parsed.status_code == 200, parsed.text
    body = parsed.json()
    assert body["source_kind"] == "xlsx_template"
    assert [item["key"] for item in body["mapped"]] == ["area_m2"]

    applied = await client.post(f"{PROJECTS}/{project_id}/import/{body['import_id']}/apply", headers=headers)
    params = {p["key"]: p for p in applied.json()["params"]}
    assert params["area_m2"]["value"] == 7000
    assert params["area_m2"]["provenance"]["status"] == "imported"
    assert params["area_m2"]["provenance"]["source"]["title"] == "Файл «my.xlsx»"


async def test_organizer_workbook_fills_blank_project(client: AsyncClient) -> None:
    project_id, headers = await _blank(client)
    parsed = (
        await client.post(
            f"{PROJECTS}/{project_id}/import",
            files={"file": (DATASET.name, DATASET.read_bytes(), XLSX)},
            headers=headers,
        )
    ).json()
    assert parsed["source_kind"] == "xlsx_freeform"
    assert len(parsed["mapped"]) == 42
    assert any("Аэропорт" in warning for warning in parsed["warnings"])

    await client.post(f"{PROJECTS}/{project_id}/import/{parsed['import_id']}/apply", headers=headers)
    report = (await client.get(f"{PROJECTS}/{project_id}/validation", headers=headers)).json()
    assert not any(issue["code"] == "REQUIRED_MISSING" for issue in report["issues"])
    processes = await client.get(f"{PROJECTS}/{project_id}/processes", headers=headers)
    assert processes.status_code == 200


async def test_user_values_are_protected_by_default(client: AsyncClient) -> None:
    project_id, headers = await _blank(client)
    await client.patch(f"{PROJECTS}/{project_id}/params/area_m2", json={"value": 5000}, headers=headers)
    parsed = (
        await client.post(
            f"{PROJECTS}/{project_id}/import",
            files={"file": (DATASET.name, DATASET.read_bytes(), XLSX)},
            headers=headers,
        )
    ).json()
    area = next(item for item in parsed["mapped"] if item["key"] == "area_m2")
    assert area["conflict_with_current"]["current_value"] == 5000
    applied = await client.post(
        f"{PROJECTS}/{project_id}/import/{parsed['import_id']}/apply", headers=headers
    )
    params = {p["key"]: p for p in applied.json()["params"]}
    assert params["area_m2"]["value"] == 5000


async def test_csv_with_semicolons(client: AsyncClient) -> None:
    project_id, headers = await _blank(client)
    csv_body = "Ключ;Значение\npickers;80\naisle_width_m;3,1\n".encode("cp1251")
    parsed = (
        await client.post(
            f"{PROJECTS}/{project_id}/import",
            files={"file": ("p.csv", csv_body, "text/csv")},
            headers=headers,
        )
    ).json()
    values = {item["key"]: item["value"] for item in parsed["mapped"]}
    assert values == {"pickers": 80, "aisle_width_m": 3.1}


async def test_unsupported_file_is_415(client: AsyncClient) -> None:
    project_id, headers = await _blank(client)
    response = await client.post(
        f"{PROJECTS}/{project_id}/import",
        files={"file": ("plan.dwg", b"AC1027", "application/octet-stream")},
        headers=headers,
    )
    assert response.status_code == 415
    assert response.json()["error_code"] == "UNSUPPORTED_FILE_TYPE"


async def test_text_import_without_llm_is_503(client: AsyncClient) -> None:
    project_id, headers = await _blank(client)
    response = await client.post(
        f"{PROJECTS}/{project_id}/import/text", json={"text": "Склад 7000 м²"}, headers=headers
    )
    assert response.status_code == 503
    assert response.json()["error_code"] == "LLM_UNAVAILABLE"
