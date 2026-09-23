import json
from collections.abc import Sequence
from io import BytesIO
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from PIL import Image
from reportlab.pdfbase.pdfmetrics import stringWidth

from app.infra.reports import pdf as pdf_module
from tests.conftest import bearer, login
from tests.integration.test_scenarios_api import _demo, _purchase

pytestmark = pytest.mark.integration


async def _report(client: AsyncClient, headers: dict[str, str], project_id: str, **payload: object) -> dict:
    response = await client.post(f"/api/v1/projects/{project_id}/reports", json=payload, headers=headers)
    assert response.status_code == 202, response.text
    # Inline jobs run after the response; the test client returns once background tasks are done.
    return (await client.get(f"/api/v1/reports/{response.json()['id']}", headers=headers)).json()


def _png() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (40, 30), "#2E86AB").save(buffer, format="PNG")
    return buffer.getvalue()


async def test_pdf_excel_and_json_reports_are_built_and_downloaded(client: AsyncClient) -> None:
    project_id, headers = await _demo(client)
    await _purchase(client, headers, project_id)

    pdf = await _report(client, headers, project_id, format="pdf")
    assert pdf["status"] == "done", pdf
    assert pdf["file"]["content_type"] == "application/pdf"
    assert pdf["versions"]["norm_set_version"]
    assert "предварительная" in pdf["disclaimer"].lower()
    download = await client.get(pdf["file"]["url"], headers=headers)
    assert download.status_code == 200
    assert download.content.startswith(b"%PDF")
    assert len(download.content) == pdf["file"]["size_bytes"]
    assert "attachment" in download.headers["content-disposition"]
    assert "filename*=UTF-8''" in download.headers["content-disposition"]

    xlsx = await _report(client, headers, project_id, format="xlsx")
    assert xlsx["status"] == "done", xlsx
    book = load_workbook(BytesIO((await client.get(xlsx["file"]["url"], headers=headers)).content))
    assert {"Сводка", "Сравнение", "Параметры объекта", "Нормативы"} <= set(book.sheetnames)
    calc = next(book[name] for name in book.sheetnames if name.startswith("Расчёт — Покупка"))
    formulas = [c.value for c in calc["C"] if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 20
    flow = next(book[name] for name in book.sheetnames if name.startswith("Поток"))
    assert str(flow["B4"].value).startswith("=J")

    data = await _report(client, headers, project_id, format="json", sections=["economics"])
    assert data["sections"] == ["economics"]
    body = json.loads((await client.get(data["file"]["url"], headers=headers)).content)
    assert [s["is_baseline"] for s in body["scenarios"]] == [True, False]
    assert body["scenarios"][1]["trace"]

    listed = (await client.get(f"/api/v1/projects/{project_id}/reports", headers=headers)).json()
    assert [r["format"] for r in listed["items"]] == ["json", "xlsx", "pdf"]


async def test_excel_without_live_formulas_keeps_the_values(client: AsyncClient) -> None:
    project_id, headers = await _demo(client)
    await _purchase(client, headers, project_id)
    report = await _report(client, headers, project_id, format="xlsx", live_formulas=False)
    book = load_workbook(BytesIO((await client.get(report["file"]["url"], headers=headers)).content))
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            assert not any(isinstance(c.value, str) and c.value.startswith("=") for c in row), sheet.title
    flow = next(book[name] for name in book.sheetnames if name.startswith("Поток"))
    assert isinstance(flow["B4"].value, float)


async def test_report_without_robotized_scenarios_fails_with_a_reason(client: AsyncClient) -> None:
    project_id, headers = await _demo(client)
    report = await _report(client, headers, project_id, format="pdf")
    assert report["status"] == "failed"
    assert "сценари" in report["error"]["detail"]
    download = await client.get(f"/api/v1/reports/{report['id']}/download", headers=headers)
    assert download.status_code == 409


async def test_docx_is_not_supported_yet(client: AsyncClient) -> None:
    project_id, headers = await _demo(client)
    response = await client.post(
        f"/api/v1/projects/{project_id}/reports", json={"format": "docx"}, headers=headers
    )
    assert response.status_code == 422


async def test_simulation_visual_goes_into_the_pdf(client: AsyncClient) -> None:
    project_id, headers = await _demo(client)
    scenario = await _purchase(client, headers, project_id)
    run = (
        await client.post(
            f"/api/v1/scenarios/{scenario['id']}/simulations",
            json={"mode": "peak", "record_events": False},
            headers=headers,
        )
    ).json()
    url = f"/api/v1/simulations/{run['id']}/visuals"
    image = _png()
    uploaded = await client.post(
        url,
        files={"file": ("frame.png", image, "image/png")},
        data={"kind": "simulation_png", "caption": "Пик: пробка у ворот"},
        headers=headers,
    )
    assert uploaded.status_code == 201, uploaded.text
    visual = uploaded.json()
    stored = await client.get(visual["file"]["url"], headers=headers)
    assert stored.content == image
    assert stored.headers["content-type"] == "image/png"
    rejected = await client.post(
        url, files={"file": ("a.txt", b"text", "text/plain")}, data={"kind": "chart_png"}, headers=headers
    )
    assert rejected.status_code == 415

    report = await _report(client, headers, project_id, format="pdf", visual_ids=[visual["id"]])
    assert report["status"] == "done", report


async def test_pdf_table_headers_fit_their_columns(
    client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A header word wider than its column is split mid-word by ReportLab («Итог / о»)."""
    tight: list[str] = []
    original = pdf_module.table

    def checked(rows: Sequence[Sequence[Any]], widths: Sequence[float], **kwargs: Any) -> Any:
        for header, width in zip(rows[0], widths, strict=True):
            room = width * pdf_module.CONTENT_WIDTH - 2 * pdf_module.CELL_PADDING
            size = pdf_module.STYLES["head"].fontSize
            tight.extend(
                f"{header}: {w}"
                for w in str(header).split()
                if stringWidth(w, pdf_module.FONT_BOLD, size) > room
            )
        return original(rows, widths, **kwargs)

    monkeypatch.setattr(pdf_module, "table", checked)
    project_id, headers = await _demo(client)
    scenario = await _purchase(client, headers, project_id)
    await client.post(
        f"/api/v1/scenarios/{scenario['id']}/simulations",
        json={"mode": "peak", "record_events": False},
        headers=headers,
    )
    report = await _report(client, headers, project_id, format="pdf")
    assert report["status"] == "done", report
    assert tight == []


async def test_foreign_user_cannot_see_reports_or_files(client: AsyncClient) -> None:
    project_id, headers = await _demo(client)
    await _purchase(client, headers, project_id)
    report = await _report(client, headers, project_id, format="json")
    stranger = bearer((await login(client, "vendor@roboscope.demo"))["access"])
    assert (await client.get(f"/api/v1/reports/{report['id']}", headers=stranger)).status_code in {403, 404}
    download = await client.get(f"/api/v1/reports/{report['id']}/download", headers=stranger)
    assert download.status_code in {403, 404}
    listed = await client.get(f"/api/v1/projects/{project_id}/reports", headers=stranger)
    assert listed.status_code in {403, 404}
