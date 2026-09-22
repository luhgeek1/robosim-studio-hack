import json
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from app.domain.reference import ObjectTypeDetail, ObjectTypeKey, ParameterDef, ParameterGroup
from app.infra.importers.matching import FUZZY_MATCH_THRESHOLD
from app.infra.importers.models import ImportFormatError, ImportSourceKind
from app.infra.importers.parser import parse_file
from app.infra.importers.template import COLUMN_KEY, COLUMN_VALUE, TEMPLATE_SHEET, build_template
from app.seeds.dataset import DATASET_FILE
from app.seeds.schemas import ObjectTypeSeed, ParameterSeed, load_object_types

REPO_ROOT = Path(__file__).resolve().parents[3]
DATASET_PATH = REPO_ROOT / DATASET_FILE


def _param(seed: ParameterSeed, group_name: str) -> ParameterDef:
    return ParameterDef(
        key=seed.key,
        name=seed.name,
        group=group_name,
        type=seed.type,
        unit=seed.unit,
        required=seed.required,
        min=seed.min,
        max=seed.max,
        step=seed.step,
        enum_values=seed.enum_values or [],
        default=None,
        hint=seed.hint,
        example=seed.example,
        affects=seed.affects,
        order=seed.order,
        dataset_row=seed.dataset_row,
    )


def _seed(key: ObjectTypeKey) -> ObjectTypeSeed:
    return next(item for item in load_object_types() if item.key == key)


def _object_type(key: ObjectTypeKey) -> ObjectTypeDetail:
    seed = _seed(key)
    groups = [
        ParameterGroup(
            key=group.key,
            name=group.name,
            order=group.order,
            parameters=[_param(param, group.name) for param in seed.parameters if param.group == group.key],
        )
        for group in seed.parameter_groups
    ]
    return ObjectTypeDetail(
        key=seed.key,
        name=seed.name,
        description=seed.description,
        industry=seed.industry,
        depth=seed.depth,
        parameter_groups=groups,
        processes=[],
        labor_groups=[],
        layout_templates=seed.layout_templates,
    )


@pytest.fixture(scope="module")
def warehouse() -> ObjectTypeDetail:
    return _object_type(ObjectTypeKey.WAREHOUSE)


def _fill_template(content: bytes, values: dict[str, object]) -> bytes:
    workbook = openpyxl.load_workbook(BytesIO(content))
    sheet = workbook[TEMPLATE_SHEET]
    header = [cell.value for cell in sheet[1]]
    key_col, value_col = header.index(COLUMN_KEY) + 1, header.index(COLUMN_VALUE) + 1
    for row in range(2, sheet.max_row + 1):
        key = sheet.cell(row, key_col).value
        if key in values:
            sheet.cell(row, value_col, values[key])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_template_layout(warehouse: ObjectTypeDetail) -> None:
    workbook = openpyxl.load_workbook(BytesIO(build_template(warehouse)))
    assert workbook.sheetnames[0] == TEMPLATE_SHEET
    assert "Как заполнить" in workbook.sheetnames
    assert workbook["_meta"].sheet_state == "hidden"
    sheet = workbook[TEMPLATE_SHEET]
    keys = [cell.value for cell in sheet["A"][1:] if cell.value]
    assert keys == [param.key for param in warehouse.parameters]
    assert sheet.freeze_panes == "B2"
    assert len(sheet.data_validations.dataValidation) >= 2


def test_template_roundtrip(warehouse: ObjectTypeDetail) -> None:
    values: dict[str, object] = {
        "area_m2": 15000,
        "floors": "2",
        "has_wms": "Да",
        "floor_type": "Эпоксидное покрытие",
        "pallet_dims_mm": "1200×800×1500",
    }
    content = _fill_template(build_template(warehouse), values)
    result = parse_file(content, "шаблон.xlsx", warehouse.parameters)
    assert result.source_kind is ImportSourceKind.XLSX_TEMPLATE
    mapped = {item.key: item for item in result.mapped}
    assert set(mapped) == set(values)
    assert mapped["area_m2"].value == 15000
    assert mapped["area_m2"].unit == "м²"
    assert mapped["floors"].value == 2
    assert mapped["has_wms"].value is True
    assert mapped["floor_type"].value == "Эпоксидное покрытие"
    assert mapped["pallet_dims_mm"].value == "1200×800×1500"
    assert all(item.confidence == 1.0 for item in result.mapped)
    assert result.unmapped == []


def test_empty_template_maps_nothing(warehouse: ObjectTypeDetail) -> None:
    result = parse_file(build_template(warehouse), "t.xlsx", warehouse.parameters)
    assert result.mapped == []
    assert result.unmapped == []


@pytest.mark.parametrize(
    ("key", "sheet"),
    [
        (ObjectTypeKey.WAREHOUSE, "Склад"),
        (ObjectTypeKey.AIRPORT, "Аэропорт"),
        (ObjectTypeKey.HOSPITAL, "Медучреждение"),
    ],
)
def test_organizer_dataset(key: ObjectTypeKey, sheet: str) -> None:
    object_type = _object_type(key)
    dataset_rows = {param.key for param in object_type.parameters if param.dataset_row}
    result = parse_file(DATASET_PATH.read_bytes(), DATASET_PATH.name, object_type.parameters)
    assert result.source_kind is ImportSourceKind.XLSX_FREEFORM
    mapped = {item.key: item for item in result.mapped}
    assert set(mapped) == dataset_rows
    assert all(item.confidence == 1.0 for item in result.mapped)
    assert result.unmapped == []
    skipped = [warning for warning in result.warnings if "не загружен" in warning]
    assert len(skipped) == 3
    assert all(f"выбран лист «{sheet}»" in warning for warning in skipped)


def test_organizer_warehouse_values(warehouse: ObjectTypeDetail) -> None:
    result = parse_file(DATASET_PATH.read_bytes(), DATASET_PATH.name, warehouse.parameters)
    mapped = {item.key: item for item in result.mapped}
    assert len(mapped) == 42
    assert mapped["area_m2"].value == 20000
    assert mapped["area_m2"].unit == "м²"
    assert mapped["area_m2"].raw_field == "Общая площадь склада"
    assert mapped["main_aisle_width_m"].value == 3.5
    assert mapped["floor_type"].value == "Промышленный бетон"


def test_csv_semicolon_cp1251(warehouse: ObjectTypeDetail) -> None:
    text = "Параметр;Ед. изм.;Значение\nОбщая площадь склада;м²;12 500\nШирина главных проездов;м;3,2\n"
    text += "Непонятная строка;;7\n"
    result = parse_file(text.encode("cp1251", errors="replace"), "data.csv", warehouse.parameters)
    assert result.source_kind is ImportSourceKind.CSV
    mapped = {item.key: item.value for item in result.mapped}
    assert mapped == {"area_m2": 12500.0, "main_aisle_width_m": 3.2}
    assert [field.raw_field for field in result.unmapped] == ["Непонятная строка"]


def test_csv_keys_utf8_bom(warehouse: ObjectTypeDetail) -> None:
    text = "key,value,unit\narea_m2,18000,м²\narea_m2,1,м²\nfloors,1,\n"
    result = parse_file(text.encode("utf-8-sig"), "data.csv", warehouse.parameters)
    assert {item.key: item.value for item in result.mapped} == {"area_m2": 18000.0, "floors": 1}
    assert any("повторяет" in warning for warning in result.warnings)


def test_json_mapping(warehouse: ObjectTypeDetail) -> None:
    payload = {"area_m2": 20000, "Высота потолков в зоне хранения": 12, "has_wms": True, "floors": None}
    result = parse_file(json.dumps(payload).encode(), "p.json", warehouse.parameters)
    assert result.source_kind is ImportSourceKind.JSON
    assert {item.key: item.value for item in result.mapped} == {
        "area_m2": 20000,
        "ceiling_height_m": 12,
        "has_wms": True,
    }


def test_json_list(warehouse: ObjectTypeDetail) -> None:
    payload = [
        {"key": "ceiling_height_m", "value": 11000, "unit": "мм"},
        {"name": "Общая площадь склада", "value": "20 000"},
        "garbage",
    ]
    result = parse_file(json.dumps(payload, ensure_ascii=False).encode(), "p.json", warehouse.parameters)
    mapped = {item.key: item for item in result.mapped}
    assert mapped["ceiling_height_m"].value == 11000
    assert mapped["ceiling_height_m"].unit == "мм"
    assert mapped["area_m2"].value == 20000.0
    assert mapped["area_m2"].raw_value == "20 000"
    assert len(result.warnings) == 1


def test_fuzzy_match(warehouse: ObjectTypeDetail) -> None:
    text = "Параметр;Значение\nОбщая плошадь склада;20000\nВысота потолка в зоне хранения;10\n"
    result = parse_file(text.encode(), "data.csv", warehouse.parameters)
    mapped = {item.key: item for item in result.mapped}
    assert set(mapped) == {"area_m2", "ceiling_height_m"}
    assert all(FUZZY_MATCH_THRESHOLD <= item.confidence < 1.0 for item in result.mapped)


def test_ambiguous_fuzzy_goes_to_unmapped_with_suggestion(warehouse: ObjectTypeDetail) -> None:
    text = "Параметр;Значение\nОбщая площадь;20000\n"
    result = parse_file(text.encode(), "data.csv", warehouse.parameters)
    assert result.mapped == []
    assert result.unmapped[0].raw_field == "Общая площадь"
    assert result.unmapped[0].suggestion_key is not None


def test_bad_value_is_a_warning(warehouse: ObjectTypeDetail) -> None:
    text = "Параметр;Значение\nОбщая площадь склада;много\nКоличество этажей (мезонинов);1\n"
    result = parse_file(text.encode(), "data.csv", warehouse.parameters)
    assert [item.key for item in result.mapped] == ["floors"]
    assert any("не число" in warning for warning in result.warnings)


@pytest.mark.parametrize("filename", ["data.pdf", "data.xls", "noextension"])
def test_unsupported_format(warehouse: ObjectTypeDetail, filename: str) -> None:
    with pytest.raises(ImportFormatError, match="Неподдерживаемый формат"):
        parse_file(b"%PDF-1.4", filename, warehouse.parameters)


def test_broken_xlsx(warehouse: ObjectTypeDetail) -> None:
    with pytest.raises(ImportFormatError):
        parse_file(b"not a zip", "broken.xlsx", warehouse.parameters)
