from dataclasses import asdict
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook
from PIL import Image

from app.domain.reports.models import ReportModel, ScenarioReport
from app.domain.scenario.models import Financing, FinancingKind
from app.engine.calculation import calculate
from app.infra.reports import pdf
from app.infra.reports.xlsx import CalculationSheet, _cashflow, _simulation
from app.service.scenarios.serialize import Serializer
from tests.unit.engine.test_calculation import scenario


def _scenario(*, down: float) -> ScenarioReport:
    calculated = calculate(scenario(financing=Financing(kind=FinancingKind.LOAN, down_payment_pct=down)))
    economics = calculated.economics
    return ScenarioReport(
        id="financed",
        name="Кредит",
        kind="purchase",
        is_baseline=False,
        is_recommended=True,
        financing={},
        calculated_at=datetime.now(UTC),
        result={
            "metrics": asdict(economics.metrics),
            "cashflow": {
                "monthly": Serializer.cashflow(economics.monthly),
                "upfront_rub": economics.upfront,
            },
            "owner_cashflow": {
                "monthly": Serializer.cashflow(economics.owner_monthly),
                "upfront_rub": economics.owner_upfront,
            },
        },
        trace=[
            {**asdict(step), "metric_key": step.key, "formula_rendered": step.rendered}
            for step in calculated.trace
        ],
        narrative={},
    )


@pytest.mark.parametrize("down", [0, 30])
def test_exported_project_and_owner_values_match_their_own_metrics(down: float) -> None:
    report = _scenario(down=down)
    book = Workbook()
    for owner, prefix, title in (
        (False, "", "Поток проекта"),
        (True, "owner_", "Поток собственника"),
    ):
        _cashflow(book, report, None, live=False, owner=owner)
        sheet = next(sheet for sheet in book.worksheets if sheet.title.startswith(title))
        metrics = report.result["metrics"]
        assert sheet["B3"].value == pytest.approx(metrics[f"{prefix}discount_rate_pct"] / 100)
        assert sheet["B4"].value == pytest.approx(metrics[f"{prefix}npv_rub"])
        roi = metrics[f"{prefix}roi_pct"]
        assert sheet["B6"].value == (pytest.approx(roi / 100) if roi is not None else None)
        data = [row for row in sheet.iter_rows() if isinstance(row[0].value, int)]
        upfront = report.result["owner_cashflow" if owner else "cashflow"]["upfront_rub"]
        assert data[0][0].value == 0
        assert data[0][5].value == -upfront
        assert sum(row[5].value for row in data) == pytest.approx(data[-1][6].value)


def test_live_excel_uses_separate_rates_and_handles_undefined_owner_roi() -> None:
    report = _scenario(down=0)
    book = Workbook()
    calc = CalculationSheet(book, report)
    calc.write()
    _cashflow(book, report, calc)
    _cashflow(book, report, calc, owner=True)
    project = next(sheet for sheet in book.worksheets if sheet.title.startswith("Поток проекта"))
    owner = next(sheet for sheet in book.worksheets if sheet.title.startswith("Поток собственника"))
    assert project["B3"].value == f"={calc.ref('discount_rate')}"
    assert owner["B3"].value == f"={calc.ref('owner_discount_rate')}"
    assert project["B3"].value != owner["B3"].value
    assert str(owner["B6"].value).startswith("=IF(F")
    assert str(owner["B5"].value).startswith("=IFERROR(")
    assert str(owner["B4"].value).startswith("=J")


def test_pdf_financial_charts_use_the_same_project_and_owner_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    report = _scenario(down=30)
    captured: list[tuple[list[dict[str, Any]], float]] = []
    buffer = BytesIO()
    Image.new("RGB", (40, 30)).save(buffer, format="PNG")

    def chart(rows: list[dict[str, Any]], upfront: float) -> bytes:
        captured.append((rows, upfront))
        return buffer.getvalue()

    monkeypatch.setattr(pdf.charts, "cashflow_chart", chart)
    pdf._fonts()
    builder = object.__new__(pdf.Builder)
    builder.story = []
    builder._financial_perspectives(report.result)
    assert captured == [
        (report.result[key]["monthly"], report.result[key]["upfront_rub"])
        for key in ("cashflow", "owner_cashflow")
    ]


def test_report_explains_missing_simulation_evidence_and_keeps_diagnostics() -> None:
    report = _scenario(down=30)
    note = "Нет актуального подтверждения парка имитацией; показана диагностика перебора."
    report.simulation = {
        "summary": None,
        "fleet": [],
        "note": note,
        "sweep": {
            "applied": False,
            "recommended_count": None,
            "points": [{"count": 1, "sla_achieved_pct": 30, "utilization": 0.9}],
            "explanation": "Целевой SLA не достигнут.",
        },
    }
    pdf._fonts()
    builder = object.__new__(pdf.Builder)
    builder.story = []
    builder._simulation(report)
    text = " ".join(item.getPlainText() for item in builder.story if isinstance(item, pdf.Paragraph))
    assert note in text
    assert "не применяется к расчёту" in text
    assert "Целевой SLA не достигнут." in text
    assert any(isinstance(item, pdf.Image) for item in builder.story)
    model = ReportModel("Оценка", {}, datetime.now(UTC), "", {}, [], scenarios=[report])
    book = Workbook()
    _simulation(book, model)
    assert note in [cell.value for row in book["Имитация"].iter_rows() for cell in row]
