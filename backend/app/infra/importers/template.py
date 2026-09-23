from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.reference import ObjectTypeDetail, ParameterDef

TEMPLATE_SHEET = "Параметры"
INSTRUCTIONS_SHEET = "Как заполнить"
META_SHEET = "_meta"
LISTS_SHEET = "_lists"
TEMPLATE_MARKER = "robomera-parameters-template"
TEMPLATE_VERSION = "1"
META_MARKER_KEY = "template"
META_OBJECT_TYPE_KEY = "object_type"
META_VERSION_KEY = "version"

COLUMN_KEY = "Ключ"
COLUMN_NAME = "Параметр"
COLUMN_GROUP = "Группа"
COLUMN_UNIT = "Ед. изм."
COLUMN_VALUE = "Значение"
COLUMN_REQUIRED = "Обязательный"
COLUMN_RANGE = "Диапазон"
COLUMN_EXAMPLE = "Пример"
COLUMN_HINT = "Подсказка"
COLUMN_WIDTHS: dict[str, int] = {
    COLUMN_KEY: 24,
    COLUMN_NAME: 48,
    COLUMN_GROUP: 26,
    COLUMN_UNIT: 10,
    COLUMN_VALUE: 22,
    COLUMN_REQUIRED: 14,
    COLUMN_RANGE: 18,
    COLUMN_EXAMPLE: 16,
    COLUMN_HINT: 80,
}
COLUMNS = list(COLUMN_WIDTHS)
HEADER_ROW = 1
YES_LABEL = "Да"
NO_LABEL = "Нет"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_GROUP_FILL = PatternFill("solid", fgColor="D9E1F2")
_REQUIRED_FILL = PatternFill("solid", fgColor="FFF2CC")
_VALUE_FILL = PatternFill("solid", fgColor="E2EFDA")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_KEY_FONT = Font(color="808080")
_INSTRUCTIONS_WIDTH = 110

_INSTRUCTIONS = [
    "Как заполнить шаблон параметров объекта",
    "",
    "1. Откройте лист «Параметры». Одна строка — один параметр объекта.",
    "2. Впишите значения только в колонку «Значение» (зелёная). Остальные колонки — справочные.",
    "3. Жёлтым выделены обязательные параметры: без них расчёт не запустится.",
    "4. Числа пишите без единиц измерения: единица указана в колонке «Ед. изм.».",
    "   Если данные в другой единице (например, мм вместо м), исправьте «Ед. изм.» — платформа пересчитает.",
    "5. Для параметров со списком выберите вариант из выпадающего списка. Для «да/нет» — «Да» или «Нет».",
    "6. Колонка «Диапазон» показывает типичные значения; выход за диапазон — повод перепроверить данные.",
    "7. Пустые значения можно оставить: платформа подставит значения по умолчанию и пометит их допущениями.",
    "8. Не меняйте колонку «Ключ» и не удаляйте скрытые листы — по ним платформа узнаёт параметры.",
    "",
    "Сохраните файл в формате .xlsx и загрузите его на шаге «Параметры объекта».",
]


def _format_number(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else f"{value:g}".replace(".", ",")


def _range_text(param: ParameterDef) -> str:
    if param.min is not None and param.max is not None:
        return f"{_format_number(param.min)} – {_format_number(param.max)}"
    if param.min is not None:
        return f"от {_format_number(param.min)}"
    if param.max is not None:
        return f"до {_format_number(param.max)}"
    return ""


def _row_cells(param: ParameterDef, group_name: str) -> list[str]:
    return [
        param.key,
        param.name,
        group_name,
        param.unit or "",
        "",
        YES_LABEL if param.required else "",
        _range_text(param),
        param.example or "",
        param.hint or "",
    ]


def _write_header(sheet: Worksheet) -> None:
    sheet.append(COLUMNS)
    for index, title in enumerate(COLUMNS, 1):
        cell = sheet.cell(HEADER_ROW, index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS[title]
    sheet.freeze_panes = sheet.cell(HEADER_ROW + 1, COLUMNS.index(COLUMN_NAME) + 1)


def _style_param_row(sheet: Worksheet, row: int, required: bool) -> None:
    sheet.cell(row, 1).font = _KEY_FONT
    for column in range(2, len(COLUMNS) + 1):
        cell = sheet.cell(row, column)
        cell.alignment = Alignment(vertical="top", wrap_text=column == len(COLUMNS))
        if required:
            cell.fill = _REQUIRED_FILL
    sheet.cell(row, COLUMNS.index(COLUMN_VALUE) + 1).fill = _VALUE_FILL
    if required:
        sheet.cell(row, COLUMNS.index(COLUMN_NAME) + 1).font = Font(bold=True)


class _Dropdowns:
    def __init__(self, workbook: Workbook, target: Worksheet) -> None:
        self._lists = workbook.create_sheet(LISTS_SHEET)
        self._lists.sheet_state = "hidden"
        self._target = target
        self._column = 0

    def add(self, param: ParameterDef, row: int) -> None:
        labels = [item["label"] for item in param.enum_values] if param.type == "enum" else []
        if param.type == "boolean":
            labels = [YES_LABEL, NO_LABEL]
        if not labels:
            return
        self._column += 1
        letter = get_column_letter(self._column)
        for index, label in enumerate(labels, 1):
            self._lists.cell(index, self._column, label)
        source = f"{quote_sheetname(LISTS_SHEET)}!${letter}$1:${letter}${len(labels)}"
        validation = DataValidation(type="list", formula1=source, allow_blank=True)
        validation.error = "Выберите значение из списка"
        validation.errorTitle = "Недопустимое значение"
        self._target.add_data_validation(validation)
        validation.add(self._target.cell(row, COLUMNS.index(COLUMN_VALUE) + 1))


def _write_parameters(workbook: Workbook, sheet: Worksheet, object_type: ObjectTypeDetail) -> None:
    dropdowns = _Dropdowns(workbook, sheet)
    for group in sorted(object_type.parameter_groups, key=lambda item: item.order):
        sheet.append(["", group.name])
        for column in range(1, len(COLUMNS) + 1):
            sheet.cell(sheet.max_row, column).fill = _GROUP_FILL
        sheet.cell(sheet.max_row, 2).font = Font(bold=True)
        for param in sorted(group.parameters, key=lambda item: item.order):
            sheet.append(_row_cells(param, group.name))
            _style_param_row(sheet, sheet.max_row, param.required)
            dropdowns.add(param, sheet.max_row)


def _write_meta(workbook: Workbook, object_type: ObjectTypeDetail) -> None:
    meta = workbook.create_sheet(META_SHEET)
    meta.sheet_state = "hidden"
    meta.append([META_MARKER_KEY, TEMPLATE_MARKER])
    meta.append([META_OBJECT_TYPE_KEY, object_type.key.value])
    meta.append([META_VERSION_KEY, TEMPLATE_VERSION])


def _write_instructions(workbook: Workbook, object_type: ObjectTypeDetail) -> None:
    sheet = workbook.create_sheet(INSTRUCTIONS_SHEET)
    sheet.column_dimensions["A"].width = _INSTRUCTIONS_WIDTH
    for line in _INSTRUCTIONS:
        sheet.append([line])
    sheet.cell(1, 1).font = Font(bold=True, size=14)
    sheet.append([])
    sheet.append([f"Тип объекта: {object_type.name}"])


def build_template(object_type: ObjectTypeDetail) -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = TEMPLATE_SHEET
    _write_header(sheet)
    _write_instructions(workbook, object_type)
    _write_parameters(workbook, sheet, object_type)
    _write_meta(workbook, object_type)
    workbook.properties.subject = f"{TEMPLATE_MARKER}:{object_type.key.value}"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
