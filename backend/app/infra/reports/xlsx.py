import re
from collections.abc import Callable
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.reports.models import ReportModel, ReportSection, ScenarioReport
from app.infra.reports.formulas import translate
from app.infra.reports.labels import KIND_LABELS, STATUS_LABELS, scenario_kind, verdict_label

SHEET_NAME_MAX = 31
MONTHS_PER_YEAR = 12
PERCENT = 100
_HEADER = PatternFill("solid", fgColor="1F3A5F")
_SECTION = PatternFill("solid", fgColor="DCE6F1")
_INPUT = PatternFill("solid", fgColor="FFF2CC")
_LIVE = PatternFill("solid", fgColor="E2EFDA")
_WARN = PatternFill("solid", fgColor="FCE4D6")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_RUB = '#,##0" ₽"'
_NUMBER = "#,##0.00"
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(book: Workbook, name: str) -> str:
    base = _BAD_SHEET_CHARS.sub(" ", name).strip()[:SHEET_NAME_MAX]
    candidate, index = base, 2
    while candidate in book.sheetnames:
        suffix = f" {index}"
        candidate = base[: SHEET_NAME_MAX - len(suffix)] + suffix
        index += 1
    return candidate


class Sheet:
    """A worksheet written top to bottom: title, notes, tables with a header row."""

    def __init__(self, book: Workbook, name: str, widths: list[int]) -> None:
        self.ws: Worksheet = book.create_sheet(_sheet_name(book, name))
        self.row = 1
        for index, width in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(index)].width = width

    def title(self, text: str) -> None:
        cell = self.ws.cell(self.row, 1, text)
        cell.font = Font(bold=True, size=14, color="1F3A5F")
        self.row += 1

    def note(self, text: str, *, fill: PatternFill | None = None, span: int = 6) -> None:
        cell = self.ws.cell(self.row, 1, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if fill is not None:
            cell.fill = fill
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=span)
        self.ws.row_dimensions[self.row].height = max(15, 15 * (len(text) // 110 + 1))
        self.row += 1

    def section(self, text: str, span: int = 6) -> None:
        self.row += 1
        for column in range(1, span + 1):
            self.ws.cell(self.row, column).fill = _SECTION
        self.ws.cell(self.row, 1, text).font = Font(bold=True)
        self.row += 1

    def header(self, labels: list[str]) -> None:
        for column, label in enumerate(labels, start=1):
            cell = self.ws.cell(self.row, column, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = _HEADER
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = _BORDER
        self.row += 1

    def line(self, values: list[Any], formats: dict[int, str] | None = None) -> int:
        for column, value in enumerate(values, start=1):
            cell = self.ws.cell(self.row, column, value)
            cell.border = _BORDER
            cell.alignment = Alignment(wrap_text=isinstance(value, str), vertical="top")
            if formats and column in formats:
                cell.number_format = formats[column]
        self.row += 1
        return self.row - 1


def _metric_row(name: str, key: str, unit: str) -> tuple[str, str, str]:
    return name, key, unit


_SUMMARY_METRICS = [
    _metric_row("CAPEX", "capex_rub", "₽"),
    _metric_row("OPEX роботизации в год", "opex_rub_year", "₽/год"),
    _metric_row("Чистый годовой эффект", "effect_rub_year", "₽/год"),
    _metric_row("Простой срок окупаемости (ТЗ)", "payback_years", "лет"),
    _metric_row("Дисконтированная окупаемость", "discounted_payback_years", "лет"),
    _metric_row("ROI за горизонт", "roi_pct", "%"),
    _metric_row("NPV", "npv_rub", "₽"),
    _metric_row("IRR", "irr_pct", "%"),
    _metric_row("TCO за горизонт", "tco_rub", "₽"),
    _metric_row("Роботов", "robots_total", "шт"),
    _metric_row("Высвобождается персонала", "fte_released", "FTE"),
]


def _summary(book: Workbook, model: ReportModel) -> None:
    sheet = Sheet(book, "Сводка", [38, 22, 22, 22, 22, 22])
    sheet.title(model.title)
    project = model.project
    generated = f"{model.generated_at:%d.%m.%Y %H:%M} UTC"
    sheet.note(
        f"Объект: {project['name']} ({project['object_type_name']}). Сформирован {generated}, {model.author}."
    )
    sheet.note(model.disclaimer, fill=_WARN)
    versions = model.versions
    sheet.note(
        "Версии: проект v{project_version}, каталог {catalog_version}, нормативы {norm_set_version}, движок "
        "{engine_version}, планировка v{layout_version}".format(
            **{
                k: versions.get(k, "—")
                for k in (
                    "project_version",
                    "catalog_version",
                    "norm_set_version",
                    "engine_version",
                    "layout_version",
                )
            }
        )
    )
    sheet.note(
        "Как пользоваться: на листах «Расчёт» жёлтые ячейки — входные данные, зелёные — живые формулы. "
        "Измените вход — пересчитаются количество роботов, CAPEX, OPEX, эффект и окупаемость."
    )
    scenarios = model.scenarios
    sheet.section("Сценарии")
    sheet.header(["Показатель", *[s.name for s in scenarios]])
    sheet.line(["Вид", *[scenario_kind(s.kind) for s in scenarios]])
    sheet.line(["Вердикт", *[verdict_label(s.result["interpretation"]["verdict"]) for s in scenarios]])
    for name, key, unit in _SUMMARY_METRICS:
        values = [s.result["metrics"].get(key) for s in scenarios]
        fmt = _RUB if unit in {"₽", "₽/год"} else _NUMBER
        sheet.line([f"{name}, {unit}", *values], dict.fromkeys(range(2, len(scenarios) + 2), fmt))
    recommended = next((s for s in scenarios if s.is_recommended), None)
    if model.comparison:
        sheet.section("Рекомендация")
        for text in model.comparison.get("recommendation", {}).get("rationale", []):
            sheet.note(f"• {text}")
    for scenario in scenarios:
        if scenario.is_baseline:
            continue
        sheet.section(f"{scenario.name}{' — рекомендован' if scenario is recommended else ''}")
        sheet.note(scenario.narrative.get("executive_summary", ""))


class CalculationSheet:
    """Inputs block (yellow, editable) above the calculation block (green: live formulas from the trace)."""

    def __init__(self, book: Workbook, scenario: ScenarioReport, *, live: bool = True) -> None:
        self.scenario = scenario
        self.live = live
        self.sheet = Sheet(book, f"Расчёт — {scenario.name}", [44, 46, 18, 12, 60, 50, 12, 10])
        self.cells: dict[str, str] = {}
        self.values: dict[str, float] = {}
        self.name = self.sheet.ws.title

    def ref(self, key: str) -> str:
        return f"'{self.name}'!{self.cells[key]}"

    def write(self) -> None:
        sheet, scenario = self.sheet, self.scenario
        sheet.title(f"Расчёт: {scenario.name} ({scenario_kind(scenario.kind)})")
        sheet.note(
            "Жёлтые ячейки — входные данные: их можно менять. Зелёные — формулы, пересчитываются сами."
            if self.live
            else "Значения платформы без формул: формула каждого шага — текстом в столбце «Формула»."
        )
        self._inputs()
        self._steps()

    def _inputs(self) -> None:
        sheet = self.sheet
        sheet.section("Входные данные", 8)
        sheet.header(["Ключ", "Название", "Значение", "Ед.", "Вид", "Происхождение", "", ""])
        seen: set[str] = set()
        for step in self.scenario.trace:
            for quantity in step["inputs"]:
                key, value = quantity["key"], quantity["value"]
                if quantity["kind"] == "metric" or key in seen or not isinstance(value, int | float):
                    continue
                seen.add(key)
                provenance = quantity.get("provenance") or {}
                source = (provenance.get("source") or {}).get("title") or provenance.get("note") or ""
                status = STATUS_LABELS.get(provenance.get("status", ""), provenance.get("status", ""))
                row = sheet.line(
                    [
                        key,
                        quantity["name"],
                        value,
                        quantity.get("unit"),
                        KIND_LABELS.get(quantity["kind"], ""),
                        f"{status}. {source}".strip(". "),
                    ],
                    {3: _NUMBER},
                )
                sheet.ws.cell(row, 3).fill = _INPUT
                self.cells[key] = f"C{row}"
                self.values[key] = float(value)

    def _resolver(self, step: dict[str, Any]) -> Any:
        inputs = {q["key"]: q for q in step["inputs"]}

        def resolve(name: str) -> tuple[str, float] | None:
            candidates = [name] + [k for k in inputs if k.endswith(f".{name}")]
            prefix = step["metric_key"].rsplit(".", 1)[0] if "." in step["metric_key"] else ""
            if prefix:
                candidates.append(f"{prefix}.{name}")
            for key in candidates:
                if key in self.cells:
                    return self.cells[key], self.values[key]
            return None

        return resolve

    def _steps(self) -> None:
        sheet = self.sheet
        sheet.section("Расчёт (трасса платформы)", 8)
        sheet.header(["Шаг", "Название", "Значение", "Ед.", "Формула", "Подстановка", "Раздел", "Живая"])
        for step in self.scenario.trace:
            key, value = step["metric_key"], float(step["value"])
            excel = translate(step["formula"], self._resolver(step), value) if self.live else None
            row = sheet.row
            sheet.line(
                [
                    key,
                    step["name"],
                    f"={excel}" if excel else value,
                    step.get("unit"),
                    step["formula"],
                    step.get("formula_rendered", ""),
                    step.get("section", ""),
                    "да" if excel else "нет",
                ],
                {3: _NUMBER},
            )
            sheet.ws.cell(row, 3).fill = _LIVE if excel else PatternFill()
            self.cells[key] = f"C{row}"
            self.values[key] = value


def _rate_ref(calc: CalculationSheet) -> str | None:
    for key in ("scenario_discount_rate", "discount_rate"):
        if key in calc.cells:
            return calc.ref(key)
    return None


def _cashflow(
    book: Workbook, scenario: ScenarioReport, calc: CalculationSheet | None, *, live: bool = True
) -> None:
    cashflow = scenario.result.get("cashflow", {})
    monthly = cashflow.get("monthly", [])
    if not monthly:
        return
    sheet = Sheet(book, f"Поток — {scenario.name}", [10, 18, 18, 18, 18, 18, 20, 14, 18, 20])
    sheet.title(f"Денежный поток относительно «как сейчас»: {scenario.name}")
    sheet.note(
        "Помесячно, с внедрением, разгоном, индексацией и заменами. "
        + (
            "Накопленный поток, дисконтирование, NPV и IRR — живые формулы Excel; "
            "ставка связана с листом расчёта."
            if live
            else "Значения платформы без формул."
        )
    )
    rate_row = sheet.line(["Ставка дисконтирования", None], {2: "0.00%"})
    rate = _rate_ref(calc) if calc else None
    npv_row = sheet.line(["NPV", None], {2: _RUB})
    irr_row = sheet.line(["IRR (годовая)", None], {2: "0.0%"})
    sheet.row += 1
    sheet.header(
        [
            "Месяц",
            "CAPEX",
            "OPEX",
            "Экономия ФОТ",
            "Финансирование",
            "Чистый поток",
            "Накопленный",
            "Коэф. дисконт.",
            "Дисконт. поток",
            "Дисконт. накопленный",
        ]
    )
    upfront = float(cashflow.get("upfront_rub") or 0.0)
    rows = [
        {
            "period": 0,
            "capex_rub": upfront,
            "opex_rub": 0.0,
            "savings_rub": 0.0,
            "financing_rub": 0.0,
            "net_rub": -upfront,
        }
    ]
    for point in monthly:
        first = point["period"] == 1
        rows.append(
            {
                **point,
                "capex_rub": point["capex_rub"] - (upfront if first else 0.0),
                "net_rub": point["net_rub"] + (upfront if first else 0.0),
            }
        )
    first_row = sheet.row
    rate_value = scenario.result["metrics"]["discount_rate_pct"] / PERCENT
    cumulative_value = discounted_value = 0.0
    for point in rows:
        r = sheet.row
        factor = (1 + rate_value) ** (-point["period"] / MONTHS_PER_YEAR)
        cumulative_value += point["net_rub"]
        discounted_value += point["net_rub"] * factor
        if live:
            formulas: list[str | float] = [
                f"=F{r}" if r == first_row else f"=G{r - 1}+F{r}",
                f"=(1+$B${rate_row})^(-A{r}/12)",
                f"=F{r}*H{r}",
                f"=I{r}" if r == first_row else f"=J{r - 1}+I{r}",
            ]
        else:
            formulas = [cumulative_value, factor, point["net_rub"] * factor, discounted_value]
        sheet.line(
            [
                point["period"],
                point["capex_rub"],
                point["opex_rub"],
                point["savings_rub"],
                point["financing_rub"],
                point["net_rub"],
                *formulas,
            ],
            {**dict.fromkeys((2, 3, 4, 5, 6, 7, 9, 10), _RUB), 8: "0.0000"},
        )
    last = sheet.row - 1
    irr = scenario.result["metrics"].get("irr_pct")
    if live:
        sheet.ws.cell(rate_row, 2, f"={rate}" if rate else rate_value)
        sheet.ws.cell(npv_row, 2, f"=J{last}")
        sheet.ws.cell(irr_row, 2, f"=(1+IRR(F{first_row}:F{last}))^12-1")
    else:
        sheet.ws.cell(rate_row, 2, rate_value)
        sheet.ws.cell(npv_row, 2, discounted_value)
        sheet.ws.cell(irr_row, 2, irr / PERCENT if irr is not None else None)


def _comparison(book: Workbook, model: ReportModel) -> None:
    comparison = model.comparison
    if not comparison:
        return
    sheet = Sheet(book, "Сравнение", [40, 22, 22, 22, 22, 22])
    sheet.title("Сравнение сценариев")
    names = comparison["names"]
    sheet.header(["Показатель", *names.values()])
    for row in comparison["rows"]:
        values = [row["values"].get(sid) for sid in names]
        unit = row.get("unit") or ""
        fmt = _RUB if unit.startswith("₽") else _NUMBER
        sheet.line(
            [f"{row['name']}{f', {unit}' if unit else ''}", *values],
            dict.fromkeys(range(2, len(names) + 2), fmt),
        )


def _params(book: Workbook, model: ReportModel) -> None:
    sheet = Sheet(book, "Параметры объекта", [30, 44, 18, 12, 22, 60])
    sheet.title("Параметры объекта и их происхождение")
    sheet.header(["Группа", "Параметр", "Значение", "Ед.", "Происхождение", "Источник"])
    for param in model.params:
        sheet.line(
            [
                param.group,
                param.name,
                param.value,
                param.unit,
                STATUS_LABELS.get(param.status, param.status),
                param.source or "",
            ]
        )


def _sensitivity(book: Workbook, model: ReportModel) -> None:
    if not (model.sensitivity or model.monte_carlo or model.survey):
        return
    sheet = Sheet(book, "Риски", [44, 18, 18, 18, 18, 18, 60])
    analysed = model.analysed
    sheet.title(f"Риски и чувствительность: {analysed.name if analysed else ''}")
    if model.sensitivity:
        sheet.section("Торнадо (NPV, ₽)", 7)
        sheet.header(["Драйвер", "Низ", "Верх", "NPV при низе", "NPV при верхе", "Размах", "Вид"])
        for item in model.sensitivity["items"]:
            sheet.line(
                [
                    item["name"],
                    item["low"],
                    item["high"],
                    item["metric_at_low"],
                    item["metric_at_high"],
                    item["swing"],
                    item["kind"],
                ],
                {4: _RUB, 5: _RUB, 6: _RUB},
            )
    if model.monte_carlo:
        mc = model.monte_carlo
        sheet.section(f"Монте-Карло: {mc['runs']} прогонов", 7)
        for label, key in (("P10", "p10"), ("P50", "p50"), ("P90", "p90"), ("Среднее", "mean")):
            sheet.line([f"NPV {label}", mc.get(key)], {2: _RUB})
        for label, value in mc.get("probabilities", {}).items():
            sheet.line([label, value], {2: "0.0%"})
    if model.survey:
        sheet.section("Что уточнить при обследовании", 7)
        sheet.header(["Параметр", "Сейчас", "Размах NPV", "", "", "", "Рекомендация"])
        for item in model.survey:
            sheet.line(
                [
                    item["name"],
                    item["current_value"],
                    item["swing"],
                    None,
                    None,
                    None,
                    item["recommendation"],
                ],
                {3: _RUB},
            )


def _fleet(run: ScenarioReport) -> str:
    fleet = run.simulation["fleet"] if run.simulation else []
    return ", ".join(f"{f['product_name']}: {f['count']}" for f in fleet)


_SIMULATION_ROWS: list[tuple[str, Callable[[dict[str, Any]], Any]]] = [
    ("SLA достигнут, %", lambda s: s["sla"]["achieved_pct"]),
    ("Цель SLA, %", lambda s: s["sla"]["target_pct"]),
    ("Задач за окно", lambda s: s["demand_total"]),
    ("Выполнено роботами", lambda s: s["completed"]),
    ("Передано людям", lambda s: s["completed_by_humans"]),
    ("Загрузка парка", lambda s: s["utilization"]["fleet"]),
    ("Средняя очередь", lambda s: s["queue"]["avg"]),
    ("Ожидание в заторах, ч", lambda s: s["congestion"]["total_wait_hours"]),
    ("Узкое место", lambda s: s["bottleneck"]["explanation"]),
    ("Расчёт против имитации", lambda s: (s.get("vs_analytic") or {}).get("text")),
]


def _simulation(book: Workbook, model: ReportModel) -> None:
    runs = [s for s in model.scenarios if s.simulation]
    if not runs:
        return
    sheet = Sheet(book, "Имитация", [44, 24, 24, 24, 24])
    sheet.title("Имитация (DES на планировке объекта)")
    sheet.header(["Показатель", *[s.name for s in runs]])
    sheet.line(["Режим", *[(r.simulation or {}).get("config", {}).get("mode") for r in runs]])
    sheet.line(["Роботов", *[_fleet(r) for r in runs]])
    summaries = [(r.simulation or {}).get("summary") or {} for r in runs]
    for label, getter in _SIMULATION_ROWS:
        sheet.line([label, *[getter(summary) if summary else None for summary in summaries]])


def _assumptions(book: Workbook, model: ReportModel) -> None:
    sheet = Sheet(book, "Нормативы", [34, 40, 14, 12, 40, 80])
    sheet.title("Нормативы и допущения, использованные в расчёте")
    sheet.header(["Ключ", "Название", "Значение", "Ед.", "Источник", "Обоснование"])
    for norm in model.assumptions:
        source = (norm.get("source") or {}).get("title", "")
        sheet.line(
            [norm["key"], norm["name"], norm["value"], norm.get("unit"), source, norm.get("rationale") or ""]
        )
    sheet = Sheet(book, "Источники", [60, 20, 60])
    sheet.title("Источники")
    sheet.header(["Источник", "Вид", "Ссылка"])
    for source in model.sources:
        sheet.line([source["title"], source.get("kind", ""), source.get("url") or ""])


def render_xlsx(model: ReportModel) -> bytes:
    book = Workbook()
    default = book.active
    _summary(book, model)
    if default is not None:
        book.remove(default)
    for scenario in model.scenarios:
        calc = None
        if model.has(ReportSection.ECONOMICS) and scenario.trace:
            calc = CalculationSheet(book, scenario, live=model.live_formulas)
            calc.write()
        if model.has(ReportSection.ECONOMICS) and not scenario.is_baseline:
            _cashflow(book, scenario, calc, live=model.live_formulas)
    if model.has(ReportSection.COMPARISON):
        _comparison(book, model)
    if model.has(ReportSection.OBJECT_PARAMS):
        _params(book, model)
    if model.has(ReportSection.SENSITIVITY) or model.has(ReportSection.MONTE_CARLO):
        _sensitivity(book, model)
    if model.has(ReportSection.SIMULATION):
        _simulation(book, model)
    if model.has(ReportSection.ASSUMPTIONS) or model.has(ReportSection.SOURCES):
        _assumptions(book, model)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
